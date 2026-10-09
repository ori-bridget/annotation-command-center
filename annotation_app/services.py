from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent.parent


def read_template(filename: str) -> str:
    path = ROOT / "templates" / filename
    return path.read_text(encoding="utf-8") if path.exists() else ""


def make_report(jobs: list[dict], issues: list[dict]) -> str:
    total = len(jobs)
    completed = sum(job["status"] == "Completed" for job in jobs)
    avg_score = round(sum(job["score"] for job in jobs) / total, 1) if total else 0
    hard_jobs = [job for job in jobs if job["hard"]]
    open_issues = [issue for issue in issues if issue["status"] == "Open"]
    hard_lines = "\n".join(f"- {job['job_id']}: điểm {job['score']}, {job['edits']} lần sửa — cần xem xét." for job in hard_jobs) or "- Chưa phát hiện ca khó."
    issue_lines = "\n".join(f"- {issue['id']} — {issue['title']} ({issue['job_id']})" for issue in open_issues) or "- Không có câu hỏi mở."
    return f"""# Báo cáo trước phiên mentor

**Batch:** {jobs[0]['batch'] if jobs else 'Chưa chọn'}  
**Tạo lúc:** {datetime.now().strftime('%d/%m/%Y %H:%M')}

## 1. Tiến độ và kết quả

- Hoàn thành **{completed}/{total} job** ({round(completed / total * 100) if total else 0}%).
- Điểm trung bình: **{avg_score}/100**.
- Dữ liệu được tổng hợp từ CVAT và GitHub issues đã lọc.

## 2. Chất lượng và ca khó

{hard_lines}

## 3. Vấn đề và câu hỏi cần mentor

{issue_lines}

## 4. Kế hoạch tiếp theo

- Mentor xác nhận các câu hỏi đang mở.
- Rà soát các job có điểm thấp hoặc nhiều lần sửa.
- Cập nhật guideline/FAQ nếu có kết luận mới.
"""


def make_weekly_report(jobs: list[dict], issues: list[dict], week: str) -> str:
    total = len(jobs)
    completed = sum(job["status"] == "Completed" for job in jobs)
    progress = round(completed / total * 100) if total else 0
    avg_score = round(sum(job["score"] for job in jobs) / total, 1) if total else 0
    open_issues = [issue for issue in issues if issue["status"] == "Open"]
    rows = []
    for index, job in enumerate(jobs, start=1):
        icon = "✅" if job["status"] == "Completed" else "⛔" if job["status"] == "Blocked" else "🟡" if job["status"] in {"In progress", "In review"} else "⬜"
        rows.append(f"| {index} | {job['job_id']} — {job['category']} | @{job['assignee'].lower()} | — | {icon} {job['score']} điểm | {job['status']} · {job['edits']} lần sửa |")
    job_rows = "\n".join(rows) or "| — | Chưa có job trong bộ lọc | — | — | ⬜ 0% | — |"
    issue_lines = "\n".join(f"- {issue['id']} ({issue['job_id']}): {issue['title']}" for issue in open_issues) or "- Không có vướng mắc đang mở."
    next_week = "\n".join(f"- Xử lý {issue['id']} và cập nhật kết luận cho {issue['job_id']}." for issue in open_issues) or "- Tiếp tục hoàn thành và review các job còn lại."
    members = ", ".join(sorted({job["assignee"] for job in jobs})) or "Chưa có dữ liệu"
    return f"""# Nhật ký tuần · {week}

**Lead tuần này:** Nhóm trưởng  
**Dữ liệu / task CVAT:** {jobs[0]['batch'] if jobs else 'Chưa chọn batch'}

## Thành viên và phân công

| Thành viên | Vị trí | Phân công tuần này |
|---|---|---|
| {members} | Annotator / Reviewer | Theo dõi các job trong batch đã chọn |

## Công việc

| # | Nội dung công việc | Annotator | Reviewer | Hoàn thành | Ghi chú |
|---|---|---|---|---|---|
{job_rows}

Mức hoàn thành hiện tại: **{progress}%** · Điểm trung bình: **{avg_score}/100**

## Tổng kết

- Đã hoàn thành: {completed} / {total} job ({progress}%).
- Job cần chú ý: {sum(job['hard'] for job in jobs)}.
- Câu hỏi GitHub đang mở: {len(open_issues)}.

## Vướng mắc

{issue_lines}

## Kế hoạch tuần tiếp theo

{next_week}
"""


def build_answer(question: str, sources: list[dict]) -> tuple[str, str]:
    if not sources:
        return ("Mình chưa tìm thấy quy định phù hợp trong guideline đã chọn. Câu hỏi cần được mentor xác nhận; hệ thống không tự suy đoán.", "needs_review")
    best = sources[0]
    return (f"Theo guideline, {best['text']}\n\nPhạm vi áp dụng: {best['category']} · {best['week']} · {best['version']}", "answered")


def safe_download_name(name: str, extension: str) -> str:
    cleaned = re.sub(r"[^\w\- ]+", "", name, flags=re.UNICODE).strip() or "mentor-report"
    return f"{cleaned}{extension}"


def report_to_docx(markdown: str) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            for part in re.split(r"(\*\*.*?\*\*)", stripped):
                run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
                run.bold = part.startswith("**") and part.endswith("**")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
